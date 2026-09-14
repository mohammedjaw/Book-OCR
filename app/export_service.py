import json, re, tempfile
from pathlib import Path
from .database import connect

def filename(title, ext):
    return re.sub(r'[<>:"/\\|?*]', '_', title).strip() + ext

def book_data(uid):
    with connect() as db:
        book=db.execute('SELECT * FROM books WHERE book_uuid=?',(uid,)).fetchone()
        pages=db.execute('SELECT * FROM book_pages WHERE book_id=? ORDER BY page_number',(book['id'],)).fetchall()
    return book,pages

def make_json(uid):
    book,pages=book_data(uid)
    data={'book':{'uuid':book['book_uuid'],'title':book['title'],'original_filename':book['original_filename'],'page_count':book['page_count']},'pages':[{'page_number':p['page_number'],'status':p['status'],'text':p['extracted_text'],'line_count':p['line_count'] or 0,'model':p['model'],'thinking_level':p['thinking_level'],'error_message':p['error_message']} for p in pages]}
    f=tempfile.NamedTemporaryFile(suffix='.json',delete=False); f.close(); Path(f.name).write_text(json.dumps(data,ensure_ascii=False,indent=2),encoding='utf-8'); return f.name,filename(book['title'],'.json')

def make_docx(uid):
    from docx import Document
    from docx.enum.text import WD_BREAK
    book,pages=book_data(uid); doc=Document(); doc.core_properties.title=book['title']
    for i,p in enumerate(pages):
        if i: doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        h=doc.add_heading(f"الصفحة {p['page_number']}",level=1); h.paragraph_format.alignment=2
        text=p['extracted_text'] if p['status']=='completed' and p['extracted_text'] is not None else '[فشل استخراج هذه الصفحة]'
        if p['status']=='failed' and p['error_message']: text += '\n'+p['error_message']
        para=doc.add_paragraph(); para.paragraph_format.alignment=2; para.add_run(text)
    f=tempfile.NamedTemporaryFile(suffix='.docx',delete=False); f.close(); doc.save(f.name); return f.name,filename(book['title'],'.docx')
